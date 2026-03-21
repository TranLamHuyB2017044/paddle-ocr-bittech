import os
import subprocess
import shutil
import platform
from src.utils.logger import Log

try:
    import docx
except ImportError:
    docx = None

def get_libreoffice_path():
    """
    Trả về lệnh chạy LibreOffice phù hợp với từng hệ điều hành
    """
    if platform.system() == "Windows":
        # Các đường dẫn mặc định của LibreOffice trên Windows
        common_paths = [
            r"C:\Program Files\LibreOffice\program\soffice.exe",
            r"C:\Program Files (x86)\LibreOffice\program\soffice.exe"
        ]
        for path in common_paths:
            if os.path.exists(path):
                return path
        return "soffice.exe" # Hy vọng nó nằm trong PATH
    else:
        # Trên Linux/Docker đã cài qua apt-get
        return "soffice"

def convert_doc_to_docx(input_path, timeout=60):
    """
    Sử dụng LibreOffice headless để convert .doc sang .docx
    """
    if not os.path.exists(input_path):
        Log.error(f"[Converter] File not found: {input_path}")
        return None

    out_dir = os.path.dirname(input_path)
    filename = os.path.basename(input_path)
    name_no_ext = os.path.splitext(filename)[0]
    expected_out = os.path.join(out_dir, f"{name_no_ext}.docx")

    # Xóa file cũ nếu đã tồn tại để tránh xung đột
    if os.path.exists(expected_out):
        try: os.remove(expected_out)
        except: pass

    soffice_path = get_libreoffice_path()
    
    # Kiểm tra xem binary có tồn tại không trước khi chạy subprocess
    if not shutil.which(soffice_path) and not os.path.exists(soffice_path):
        Log.error(f"[Converter] ❌ LibreOffice (soffice) not found at: {soffice_path}")
        Log.error("[Converter] 👉 Please install LibreOffice OR run this app via Docker on VPS.")
        return None

    try:
        Log.info(f"[Converter] Converting {filename} to .docx via {soffice_path}...")
        
        command = [
            soffice_path,
            "--headless",
            "--convert-to", "docx",
            "--outdir", out_dir,
            input_path
        ]
        
        # Chạy command
        result = subprocess.run(
            command, 
            capture_output=True, 
            text=True, 
            timeout=timeout,
            shell=(platform.system() == "Windows") # Cần shell=True trên Windows nếu dùng lệnh alias
        )
        
        if os.path.exists(expected_out):
            Log.info(f"[Converter] ✅ Success: {expected_out}")
            return expected_out
        else:
            Log.error(f"[Converter] ❌ Conversion failed. Output file not found. Stdout: {result.stdout}, Stderr: {result.stderr}")
            return None
            
    except Exception as e:
        Log.error(f"[Converter] ❌ Unexpected error: {e}")
        if "WinError 2" in str(e):
            Log.error("[Converter] 👉 Mẹo: Bạn cần cài đặt LibreOffice trên máy Local để test tính năng này ngoài Docker.")
        return None

def extract_text_from_word(file_path):
    """
    Trích xuất văn bản từ file .doc hoặc .docx.
    """
    ext = file_path.lower().split('.')[-1]
    docx_path = file_path
    is_temporary_docx = False

    if ext == 'doc':
        converted = convert_doc_to_docx(file_path)
        if not converted:
            Log.error(f"[Word] ❌ Could not convert .doc to .docx: {file_path}")
            return ""
        docx_path = converted
        is_temporary_docx = True

    full_text = ""
    if docx is not None:
        try:
            doc = docx.Document(docx_path)
            full_text = "\n".join([para.text for para in doc.paragraphs])
            
            # Cleanup nếu là file tạm
            if is_temporary_docx and os.path.exists(docx_path):
                try: os.remove(docx_path)
                except: pass
                
            return full_text.strip()
        except Exception as e:
            Log.error(f"[Word] ❌ Error reading docx: {e}")
            
    return ""
def extract_lines_from_word(file_path):
    """
    Trích xuất văn bản từ file .doc hoặc .docx theo dạng danh sách các dòng.
    Returns: List of dicts [{"page": 1, "line": int, "bbox": None, "text": str}]
    """
    ext = file_path.lower().split('.')[-1]
    docx_path = file_path
    is_temporary_docx = False

    if ext == 'doc':
        converted = convert_doc_to_docx(file_path)
        if not converted:
            Log.error(f"[Word] ❌ Could not convert .doc to .docx: {file_path}")
            return []
        docx_path = converted
        is_temporary_docx = True

    lines = []
    if docx is not None:
        try:
            doc = docx.Document(docx_path)
            line_idx = 1
            for para in doc.paragraphs:
                if para.text.strip():
                    lines.append({
                        "page": 1,
                        "line": line_idx,
                        "bbox": None,
                        "text": para.text.strip()
                    })
                    line_idx += 1
            
            # Xử lý các bảng nếu có
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            lines.append({
                                "page": 1,
                                "line": line_idx,
                                "bbox": None,
                                "text": cell.text.strip()
                            })
                            line_idx += 1
            
            # Cleanup nếu là file tạm
            if is_temporary_docx and os.path.exists(docx_path):
                try: os.remove(docx_path)
                except: pass
                
            return lines
        except Exception as e:
            Log.error(f"[Word] ❌ Error reading docx lines: {e}")
            
    return []
